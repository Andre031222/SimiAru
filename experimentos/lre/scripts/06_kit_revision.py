import random

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from comun import REVISION, lexico_actual, resumir

SEMILLA = 2027
OPCIONES_GLOSA = ["sí", "parcial", "no"]
OPCIONES_VARIEDAD = {"quechua": ["Collao", "Chanka", "ambas", "no sé"], "aimara": ["norteño", "sureño", "ambos", "no sé"]}
PREGUNTA_VARIEDAD = {"quechua": "¿En qué variedad la usa? (Collao / Chanka / ambas / no sé)",
                     "aimara": "¿En qué variedad la usa? (norteño / sureño / ambos / no sé)"}


def clase(r) -> str:
    e = r.variedad_evidencia or ""
    if r.lengua == "quechua":
        if r.variedad == "comun":
            return "comun-ambas" if e.startswith("ambas") else "comun-neutra"
        return r.variedad
    if r.variedad == "comun":
        return "comun-ambas" if e.startswith("ambas") else "comun-sin_atestiguar"
    return r.variedad


def repartir(n: int, tamanos: dict[str, int]) -> dict[str, int]:
    cuota = {k: 0 for k in tamanos}
    restante = n
    abiertos = sorted(k for k, v in tamanos.items() if v > 0)
    while restante > 0 and abiertos:
        parte = max(1, restante // len(abiertos))
        for k in list(abiertos):
            if restante == 0:
                break
            toma = min(parte, tamanos[k] - cuota[k], restante)
            cuota[k] += toma
            restante -= toma
            if cuota[k] >= tamanos[k]:
                abiertos.remove(k)
    return cuota


def muestrear(df: pd.DataFrame, plan: dict[str, int | None], rng: random.Random) -> pd.DataFrame:
    partes = []
    for c, n in plan.items():
        g = df[df.clase == c]
        n = len(g) if n is None else n
        cuota = repartir(n, g.fuente_id.value_counts().to_dict())
        for fuente, k in sorted(cuota.items()):
            h = g[g.fuente_id == fuente].sort_values("id")
            ids = rng.sample(list(h.id), k)
            partes.append(h[h.id.isin(ids)].assign(estrato_poblacion=len(h), estrato_muestra=k))
    return pd.concat(partes)


def hoja(muestra: pd.DataFrame, lengua: str, ruta):
    wb = Workbook()
    ws = wb.active
    ws.title = "revision"
    columnas = ["id", "forma", "glosa_es", "¿La glosa es correcta? (sí / parcial / no)", PREGUNTA_VARIEDAD[lengua], "comentario"]
    ws.append(columnas)
    for r in muestra.itertuples():
        ws.append([r.codigo, r.forma, r.glosa_es, None, None, None])
    gris = PatternFill("solid", fgColor="DDDDDD")
    for celda in ws[1]:
        celda.font = Font(bold=True)
        celda.fill = gris
        celda.alignment = Alignment(wrap_text=True, vertical="top")
    for letra, ancho in zip("ABCDEF", (9, 26, 44, 20, 26, 40)):
        ws.column_dimensions[letra].width = ancho
    ws.freeze_panes = "B2"
    fin = len(muestra) + 1
    v1 = DataValidation(type="list", formula1='"' + ",".join(OPCIONES_GLOSA) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Valor no válido", error="Elija sí, parcial o no.")
    v2 = DataValidation(type="list", formula1='"' + ",".join(OPCIONES_VARIEDAD[lengua]) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Valor no válido", error="Elija una opción de la lista.")
    ws.add_data_validation(v1)
    ws.add_data_validation(v2)
    v1.add(f"D2:D{fin}")
    v2.add(f"E2:E{fin}")
    for fila in ws.iter_rows(min_row=2, max_row=fin):
        for celda in fila:
            celda.alignment = Alignment(wrap_text=True, vertical="top")
    info = wb.create_sheet("datos del revisor")
    for fila in (["Nombre o código del revisor", ""], ["Lengua materna", ""], ["Variedad que habla", ""], ["Lugar de origen (provincia)", ""],
                 ["Fecha de revisión", ""], ["Lea antes la rúbrica (rubrica.md) y el consentimiento (consentimiento.md).", ""]):
        info.append(fila)
    info.column_dimensions["A"].width = 60
    info.column_dimensions["B"].width = 40
    wb.save(ruta)


def main():
    rng = random.Random(SEMILLA)
    lex = lexico_actual()
    lex["apta"] = lex.apta.fillna(False).astype(bool)
    lex["clase"] = [clase(r) for r in lex.itertuples()]
    q = lex[(lex.lengua == "quechua") & lex.apta]
    a = lex[lex.lengua == "aimara"]
    a = a[a.variedad.isin(["norte", "sur"]) | a.apta]
    n_extremos = int(a.variedad.isin(["norte", "sur"]).sum())
    resto = 100 - n_extremos
    mq = muestrear(q, {"qullaw": 50, "chanka": 50, "comun-ambas": 50, "comun-neutra": 50}, rng)
    ma = muestrear(a, {"norte": None, "sur": None, "comun-ambas": resto // 2, "comun-sin_atestiguar": resto - resto // 2}, rng)
    REVISION.mkdir(parents=True, exist_ok=True)
    claves = []
    for lengua, m, prefijo, nombre in (("quechua", mq, "Q", "hoja-quechua.xlsx"), ("aimara", ma, "A", "hoja-aimara.xlsx")):
        m = m.sample(frac=1, random_state=SEMILLA).reset_index(drop=True)
        m["codigo"] = [f"{prefijo}{i + 1:03d}" for i in range(len(m))]
        hoja(m, lengua, REVISION / nombre)
        claves.append(m[["codigo", "id", "lengua", "forma", "glosa_es", "fuente_id", "variedad", "variedad_evidencia", "clase",
                         "estrato_poblacion", "estrato_muestra"]].rename(columns={"id": "lexico_id", "variedad": "variedad_predicha"}))
    clave = pd.concat(claves)
    clave.to_csv(REVISION / "clave-oculta.csv", index=False)
    resumen = {"semilla": SEMILLA,
               "quechua": clave[clave.lengua == "quechua"].groupby(["clase", "fuente_id"]).size().rename("n").reset_index().to_dict("records"),
               "aimara": clave[clave.lengua == "aimara"].groupby(["clase", "fuente_id"]).size().rename("n").reset_index().to_dict("records"),
               "total": {"quechua": int((clave.lengua == "quechua").sum()), "aimara": int((clave.lengua == "aimara").sum())},
               "archivos": ["revision/hoja-quechua.xlsx", "revision/hoja-aimara.xlsx", "revision/clave-oculta.csv"]}
    resumir("6_kit_revision", resumen)
    print(clave.groupby(["lengua", "clase", "fuente_id"]).size().to_string())


if __name__ == "__main__":
    main()
