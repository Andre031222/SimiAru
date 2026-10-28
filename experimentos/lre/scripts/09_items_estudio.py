import json
from pathlib import Path

from openpyxl import load_workbook

LRE = Path(__file__).resolve().parents[1]
REVISION = LRE / "revision"
SALIDA = LRE.parents[1] / "api" / "recursos" / "estudio-items.json"
HOJAS = {"quechua": REVISION / "hoja-quechua.xlsx", "aimara": REVISION / "hoja-aimara.xlsx"}


def leer(lengua: str, ruta: Path) -> list[dict]:
    hoja = load_workbook(ruta, read_only=True)["revision"]
    items = []
    for codigo, forma, glosa, *_ in hoja.iter_rows(min_row=2, values_only=True):
        if not codigo:
            continue
        items.append({"codigo": str(codigo).strip(), "lengua": lengua, "forma": str(forma).strip(), "glosa_es": str(glosa).strip()})
    return items


def main():
    items = [i for lengua, ruta in HOJAS.items() for i in leer(lengua, ruta)]
    codigos = [i["codigo"] for i in items]
    if len(set(codigos)) != len(codigos):
        raise SystemExit("códigos repetidos en las hojas")
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    SALIDA.write_text(json.dumps(items, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"{len(items)} ítems en {SALIDA}")


if __name__ == "__main__":
    main()
